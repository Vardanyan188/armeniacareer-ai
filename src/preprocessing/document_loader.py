# src/preprocessing/document_loader.py
#
# Document loading layer for the MVP pipeline.
#
# Responsibility:
#   Turn raw input files into plain text / plain dicts. Nothing more.
#   - Resumes (.pdf / .docx / .txt / .md) → raw text string
#   - Job descriptions (.json)            → dict, plus a flattened raw-text view
#
# Explicit non-responsibilities (enforced by later phases, NOT here):
#   - No LLM calls.
#   - No entity parsing (no CVEntities / JDEntities construction).
#   - No PII masking (that is the input guardrail's job, downstream).
#   - No mutation of anything under data/raw.
#
# Heavy parsers (pypdf, python-docx) are imported lazily inside the functions
# that need them, so importing this module never requires those packages to be
# installed — only actually loading a PDF/DOCX does.

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import List, Union

logger = logging.getLogger(__name__)

PathLike = Union[str, Path]

# Supported input formats.
RESUME_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
JD_EXTENSIONS = {".json"}

# Default data locations (current repository layout).
DEFAULT_RESUME_DIR = "data/raw/resumes"
DEFAULT_JD_DIR = "data/raw/job_descriptions"


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------

def load_text_file(path: PathLike) -> str:
    """
    Reads a UTF-8 text file (.txt / .md) and returns its content.
    Falls back to a lenient decode if the file is not valid UTF-8.
    """
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"Text file not found: {p}")
    try:
        return p.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        logger.warning("load_text_file: %s is not valid UTF-8; decoding with errors='replace'.", p)
        return p.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_pdf_text(path: PathLike) -> str:
    """
    Extracts text from a PDF using pypdf. Returns concatenated page text.
    Pages that yield no extractable text contribute an empty string.
    """
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"PDF file not found: {p}")
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError(
            "extract_pdf_text requires the 'pypdf' package. Install it via "
            "`pip install -r requirements.txt`."
        ) from exc

    reader = PdfReader(str(p))
    pages: List[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extract_docx_text(path: PathLike) -> str:
    """
    Extracts text from a .docx using python-docx. Includes paragraph text and
    table cell text (CVs frequently lay out content in tables).
    """
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"DOCX file not found: {p}")
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError(
            "extract_docx_text requires the 'python-docx' package. Install it via "
            "`pip install -r requirements.txt`."
        ) from exc

    document = docx.Document(str(p))
    parts: List[str] = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(part for part in parts if part is not None).strip()


# ---------------------------------------------------------------------------
# Resume dispatch
# ---------------------------------------------------------------------------

def load_resume_text(path: PathLike) -> str:
    """
    Loads a resume file as raw text, dispatching on file extension.
    Supported: .pdf, .docx, .txt, .md.
    Raises ValueError for any other extension.
    """
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"Resume file not found: {p}")

    ext = p.suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(p)
    if ext == ".docx":
        return extract_docx_text(p)
    if ext in (".txt", ".md"):
        return load_text_file(p)

    raise ValueError(
        f"Unsupported resume format '{ext}' for file {p.name}. "
        f"Supported formats: {sorted(RESUME_EXTENSIONS)}."
    )


# ---------------------------------------------------------------------------
# Job description (JSON)
# ---------------------------------------------------------------------------

def load_jd_json(path: PathLike) -> dict:
    """
    Loads a job-description .json file and returns the parsed dict.
    Raises ValueError for non-.json files, and ValueError if the JSON
    top level is not an object.
    """
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"JD file not found: {p}")

    ext = p.suffix.lower()
    if ext not in JD_EXTENSIONS:
        raise ValueError(
            f"Unsupported job-description format '{ext}' for file {p.name}. "
            f"Supported formats: {sorted(JD_EXTENSIONS)}."
        )

    with p.open(encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        raise ValueError(
            f"JD JSON {p.name} must contain a top-level object, got {type(data).__name__}."
        )
    return data


# Fields used to compose a fallback raw-text view when `raw_text` is absent.
_JD_FALLBACK_FIELDS = [
    ("role_title", "Role"),
    ("company_name", "Company"),
    ("industry", "Industry"),
    ("geography", "Geography"),
    ("required_experience_years", "Required experience (years)"),
    ("required_education_level", "Required education level"),
]


def jd_json_to_raw_text(jd: dict) -> str:
    """
    Returns a plain-text view of a job description for downstream text processing.

    Preference order:
      1. The verbatim `raw_text` field if present and non-empty.
      2. A deterministic, LLM-free composition of available metadata fields.

    No parsing, normalization, or inference beyond simple field concatenation.
    """
    if not isinstance(jd, dict):
        raise ValueError(f"jd_json_to_raw_text expects a dict, got {type(jd).__name__}.")

    raw = jd.get("raw_text")
    if isinstance(raw, str) and raw.strip():
        return raw

    lines: List[str] = []
    for key, label in _JD_FALLBACK_FIELDS:
        value = jd.get(key)
        if value not in (None, "", []):
            lines.append(f"{label}: {value}")

    if not lines:
        logger.warning("jd_json_to_raw_text: no raw_text and no recognised metadata fields found.")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Directory listing
# ---------------------------------------------------------------------------

# Path segments that demo loaders must NEVER return (privacy boundary):
# real private CVs/JDs and the consent-gated candidate pool.
_EXCLUDED_DIR_SEGMENTS = {"private", "uploads", "candidate_pool"}


def _is_excluded(path: Path) -> bool:
    parts = {p.lower() for p in path.parts}
    return bool(parts & _EXCLUDED_DIR_SEGMENTS)


def _list_supported_files(
    base_dir: PathLike, extensions: set, recursive: bool = True,
) -> List[Path]:
    """
    Returns a sorted list of files under base_dir whose suffix is in `extensions`.

    Recurses into subfolders by default (files now live in category subfolders),
    but always excludes private / uploads / candidate_pool locations so real
    personal data never appears in demo selectors. Returns an empty list if the
    directory does not exist.
    """
    base = Path(base_dir)
    if not base.is_dir():
        logger.warning("_list_supported_files: directory does not exist: %s", base)
        return []
    candidates = base.rglob("*") if recursive else base.iterdir()
    matches = [
        f for f in candidates
        if f.is_file() and f.suffix.lower() in extensions and not _is_excluded(f)
    ]
    return sorted(matches)


def list_resume_files(base_dir: PathLike = DEFAULT_RESUME_DIR) -> List[Path]:
    """
    Lists supported resume files (.pdf/.docx/.txt/.md) under base_dir, recursively.
    Never returns files from private/uploads/candidate_pool locations.
    """
    return _list_supported_files(base_dir, RESUME_EXTENSIONS)


def list_jd_files(base_dir: PathLike = DEFAULT_JD_DIR) -> List[Path]:
    """
    Lists job-description files (.json) under base_dir, recursively.
    Never returns files from private/uploads/candidate_pool locations.
    """
    return _list_supported_files(base_dir, JD_EXTENSIONS)
